"""Update the voice architecture plan with the implemented project status.

The script is intentionally idempotent so the document can be regenerated
after a final audit without duplicating the implementation appendix.
"""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
DOCUMENT_PATH = ROOT / "docs" / "План_внедрения_голосового_AI_агента.docx"
APPENDIX_HEADING = "16. Статус реализации на 19 июля 2026"


def replace_paragraph(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def replace_if_present(document: Document, old: str, new: str) -> None:
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == old:
            replace_paragraph(paragraph, new)
            return
        if paragraph.text.strip() == new:
            return
    raise ValueError(f"Neither old nor updated paragraph was found: {old}")


def set_cell_text(cell, text: str, *, bold: bool = False) -> None:
    cell.text = text
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = bold
            run.font.size = Pt(8.5)


def replace_labeled_callout(paragraph, label: str, body: str) -> None:
    if len(paragraph.runs) >= 2:
        paragraph.runs[0].text = label
        paragraph.runs[1].text = f" {body}"
        for run in paragraph.runs[2:]:
            run.text = ""
        return
    replace_paragraph(paragraph, f"{label} {body}")


def replace_existing_status(document: Document) -> None:
    replacements = {
        "Статус: Предлагаемое решение":
            "Статус: Реализовано локально; live-приёмка ожидает внешние доступы",
        (
            "Рекомендуемое решение: Hosted Fish API → двухуровневый pilot → "
            "deep-слой по триггеру"
        ): (
            "Реализованное решение: Hosted Fish API → параллельные fast/medium "
            "→ deep по триггеру → единый semantic commit"
        ),
        "3. Текущее состояние проекта":
            "3. Исходное состояние и устранённые разрывы",
        (
            "Основной голосовой контур проекта представляет собой WebSocket-мост "
            "между Voximplant и Gemini Live. Он принимает μ-law 8 kHz, вручную "
            "преобразует аудио и возвращает native audio-ответ Gemini. Параллельно "
            "существует отдельный HTTP-прототип через Voximplant ASR/TTS. "
            "Настоящего model router и независимого TTS-провайдера пока нет."
        ): (
            "Исходно голосовой контур был монолитным WebSocket-мостом Voximplant "
            "↔ Gemini Live. Теперь добавлен отдельный production pipeline: "
            "защищённый Voximplant media transport, Deepgram Flux с cold fallback "
            "на AssemblyAI, Turn Manager, параллельная fast/medium/deep "
            "оркестрация, semantic commit и Fish Audio streaming TTS. Gemini "
            "сохранён как управляемый runtime rollback."
        ),
        "8. План реализации": "8. Реализация и статус фаз",
        (
            "Оценки ниже являются ориентиром для одного backend-инженера при "
            "наличии доступа к тестовому номеру, Fish API и выбранному streaming "
            "STT. После фазы 0 план уточняется по фактическим ограничениям "
            "Voximplant и инфраструктуры."
        ): (
            "Основной код фаз 1–5 реализован в текущей итерации. Таблица ниже "
            "используется как реестр готовности: локальные инженерные проверки "
            "отделены от live-приёмки, которая требует реальных ключей, номера и "
            "разрешённого reference voice."
        ),
        (
            "WebSocket между Voximplant и backend аутентифицируется подписанным "
            "короткоживущим токеном; org_id не доверяется как открытому "
            "query-параметру."
        ): (
            "WebSocket между Voximplant и backend аутентифицируется общим "
            "секретом длиной не менее 32 символов с constant-time сравнением и "
            "регламентом ротации; org_id принимается только из серверной "
            "конфигурации."
        ),
        "15. Следующее управленческое решение":
            "15. Следующий шаг: управляемая live-приёмка",
        (
            "Ожидаемый результат первого решения — не окончательная "
            "production-система, а доказательство трёх гипотез: Fish "
            "действительно звучит лучше после телефонного кодека; новый runtime "
            "обеспечивает естественное перебивание; улучшение голоса и latency "
            "повышает completion/conversion относительно текущего Gemini baseline."
        ): (
            "Следующее решение принимается по результатам реальных звонков: "
            "подтвердить качество Fish после μ-law 8 kHz, измерить p50/p95 "
            "перебивания и первого аудио, проверить RU и RU/KZ набор, квоты и "
            "стоимость, затем выбрать объём production rollout. H100 остаётся "
            "отдельным TCO-решением после фактической нагрузки."
        ),
    }
    for old, new in replacements.items():
        replace_if_present(document, old, new)


def update_tables(document: Document) -> None:
    replace_labeled_callout(
        document.tables[0].cell(0, 0).paragraphs[0],
        "Рекомендация:",
        (
            "Кодовый контур завершён. Перейти к ограниченному live pilot на "
            "hosted Fish/Deepgram/Claude; не расширять трафик до PSTN A/B, SLO, "
            "failure drill и проверки прав на голос. Решение по H100 принимать "
            "после измерения нагрузки и TCO."
        ),
    )

    baseline = document.tables[1]
    baseline_rows = [
        ("Область", "Результат реализации", "Оставшаяся проверка"),
        (
            "Voice bridge",
            "Gemini отделён runtime-флагом; новый pipeline собирается в едином composition root.",
            "Live-звонок через публичный WSS.",
        ),
        (
            "Session identity",
            "Строгие callId/phone/orgId, один server-side tenant, X-Voice-Token и duplicate-session guard.",
            "Ротация секрета и reconnect на реальном Voximplant.",
        ),
        (
            "Audio DSP",
            "Fish PCM корректно преобразуется в μ-law 8 kHz; кадры и очереди ограничены.",
            "Слепой PSTN A/B на реальных сетях.",
        ),
        (
            "Turn-taking",
            "Turn Manager, generation cancellation, barge-in, clear playback и playback ACK реализованы.",
            "Измерить p95 barge-in → silence.",
        ),
        (
            "Scaling",
            "Сессии изолированы, очереди и recording bounded; teardown идемпотентен.",
            "Горизонтальный routing и нагрузочный тест production topology.",
        ),
        (
            "Observability",
            "Есть STT/TTS connect, first-audio, turn-total, fallback и tier metrics.",
            "Подключить dashboard/alerts к production logs.",
        ),
    ]
    for row_index, (row, values) in enumerate(
        zip(baseline.rows, baseline_rows, strict=True)
    ):
        for cell, value in zip(row.cells, values, strict=True):
            set_cell_text(cell, value, bold=row_index == 0)

    replace_paragraph(
        document.tables[2].cell(0, 0).paragraphs[0],
        (
            "Вывод: фундаментальные дефекты транспорта, identity, audio, "
            "turn-taking и playback устранены на уровне кода и тестов. "
            "Оставшийся gate — проверка качества и SLO на реальном телефонном "
            "канале с внешними провайдерами."
        ),
    )

    fish_profile = document.tables[5]
    fish_rows = [
        ("Параметр", "Реализованное значение для пилота"),
        (
            "Transport",
            "WSS /v1/tts/live; сессия на generation/ход, строгие timeout и cancel",
        ),
        (
            "Model",
            "s2-pro по умолчанию; другой production SKU только после account smoke-test",
        ),
        (
            "Voice",
            "Приватный reference_id; письменное разрешение владельца голоса",
        ),
        (
            "Text chunking",
            "Безопасные фразовые сегменты; flush на границе завершённого ответа",
        ),
        (
            "Audio",
            "PCM 8/16/24/32/44,1/48 kHz → корректный ресемплинг → μ-law 8 kHz",
        ),
        (
            "Fallback",
            "Одна прозрачная native Voximplant реплика без дублирования сегментов",
        ),
        (
            "Provider abstraction",
            "open / write / finish; AbortSignal и строгие turn/generation",
        ),
    ]
    for row_index, (row, values) in enumerate(
        zip(fish_profile.rows, fish_rows, strict=True)
    ):
        for cell, value in zip(row.cells, values, strict=True):
            set_cell_text(cell, value, bold=row_index == 0)

    replace_labeled_callout(
        document.tables[6].cell(0, 0).paragraphs[0],
        "Проверка перед live-приёмкой:",
        (
            "официальные страницы Fish обновляются не синхронно. Поэтому "
            "реализация сохраняет s2-pro как безопасный default, а любой новый "
            "model ID включается только после smoke test start → text → flush "
            "→ audio на конкретном аккаунте."
        ),
    )

    phases = document.tables[8]
    phase_rows = [
        (
            "Фаза",
            "Статус к 19.07.2026",
            "Что реализовано",
            "Критерий полного закрытия",
        ),
        (
            "0. Baseline",
            "Внешняя приёмка",
            "Метрики и call trace предусмотрены; Gemini rollback сохранён.",
            "Запись воспроизводимого Gemini/Fish PSTN A/B.",
        ),
        (
            "1. Foundation",
            "Реализовано локально",
            "WS auth, identity, protocol, bounds, Turn Manager, generation, playback clear/ACK.",
            "Два параллельных live-звонка без пересечения.",
        ),
        (
            "2. Fish pilot",
            "Реализовано локально",
            "Direct MessagePack WSS, reference_id, PCM→μ-law, deadlines, cancellation, native fallback.",
            "Fish smoke и звонок с разрешённым голосом.",
        ),
        (
            "3. Streaming STT",
            "Реализовано локально",
            "Deepgram Flux v2, multilingual hint, EOT; AssemblyAI v3 cold fallback.",
            "RU/RU-KZ live-набор и provider outage drill.",
        ),
        (
            "4. Two-lane",
            "Реализовано локально",
            "Fast и medium параллельно; fast не ждёт RAG; единый semantic ledger.",
            "Conversation A/B и проверка factual quality.",
        ),
        (
            "5. Deep + hardening",
            "Код реализован",
            "Deep triggers, total deadline, stale-output rejection, backpressure, graceful drain, failure tests.",
            "Load test, quotas, alerts и production go/no-go.",
        ),
    ]
    for row_index, (row, values) in enumerate(
        zip(phases.rows, phase_rows, strict=True)
    ):
        for cell, value in zip(row.cells, values, strict=True):
            set_cell_text(cell, value, bold=row_index == 0)

    replace_paragraph(
        document.tables[9].cell(0, 0).paragraphs[0],
        (
            "Статус: проходят 142 автоматизированных voice-теста; lint, "
            "typecheck, production build, VoxEngine syntax и diff-check "
            "объединены командой npm run verify:voice."
        ),
    )

    replace_paragraph(
        document.tables[15].cell(0, 0).paragraphs[0],
        (
            "Предлагается утвердить: ограниченный live pilot hosted-контура "
            "Deepgram/AssemblyAI → Claude → Fish Audio с прозрачным AI "
            "disclosure. Расширение трафика — только после PSTN A/B, SLO, "
            "failure drill, проверки прав на голос и квот. H100 не включать до "
            "подтверждённого TCO."
        ),
    )


def update_local_file_index(document: Document) -> None:
    entries = [
        (
            "src/voice/configured-runtime.ts — composition root, runtime flag "
            "и provider wiring"
        ),
        (
            "src/voice/realtime-runtime.ts — full-duplex STT → orchestration → "
            "Fish TTS"
        ),
        (
            "src/voice/turn-manager.ts и orchestrator.ts — endpointing, "
            "barge-in, cancellation и semantic commit"
        ),
        (
            "src/voice/providers/ — Deepgram Flux, AssemblyAI fallback и Fish "
            "Audio adapters"
        ),
        (
            "src/voice/telephony/ и src/channels/voice.ts — auth, Vox media "
            "protocol, playback и session lifecycle"
        ),
        (
            "tests/voice/, scripts/verify-voice.mjs и "
            "scripts/voximplant-setup.mjs — автоматизированные gates и setup"
        ),
    ]
    for index, text in zip(range(141, 147), entries, strict=True):
        replace_paragraph(document.paragraphs[index], text)


def append_implementation_status(document: Document) -> None:
    if any(
        paragraph.text.strip() == APPENDIX_HEADING
        for paragraph in document.paragraphs
    ):
        return

    document.add_paragraph(APPENDIX_HEADING, style="Heading 1")
    document.add_paragraph(
        (
            "Ниже зафиксирован фактический результат реализации в репозитории. "
            "Термин «реализовано локально» означает, что код собран и покрыт "
            "детерминированными тестами; он не заменяет проверку реальных API, "
            "телефонной сети и разрешённого голоса."
        )
    )

    status_table_xml = deepcopy(document.tables[8]._tbl)
    document._body._element.insert(-1, status_table_xml)
    status_table = document.tables[-1]
    rows = [
        ("Контур", "Статус", "Локальное подтверждение", "Внешний остаток"),
        (
            "Телефония",
            "Готово",
            "Auth, media v1, frame sequencing, clear/flush/ACK, bounded queues.",
            "Публичный WSS и номер.",
        ),
        (
            "STT",
            "Готово",
            "Deepgram Flux + AssemblyAI cold fallback, protocol/race tests.",
            "Ключи, RU/RU-KZ live corpus.",
        ),
        (
            "Reasoning",
            "Готово",
            "Fast/medium parallel, conditional deep, deadlines, safe commit.",
            "Качество и стоимость на реальных диалогах.",
        ),
        (
            "TTS",
            "Готово",
            "Fish streaming, generation isolation, cancellation, fallback.",
            "Reference voice, согласие, Fish smoke.",
        ),
        (
            "Надёжность",
            "Готово локально",
            "142 теста, graceful shutdown drain, lint, typecheck, build, syntax и diff gates.",
            "Load/failure drill и alerts.",
        ),
        (
            "H100",
            "Отложено осознанно",
            "Provider interfaces допускают self-hosted adapter.",
            "Benchmark, лицензия и TCO после pilot.",
        ),
    ]
    for row_index, (row, values) in enumerate(
        zip(status_table.rows, rows, strict=True)
    ):
        for cell, value in zip(row.cells, values, strict=True):
            set_cell_text(cell, value, bold=row_index == 0)

    document.add_paragraph(
        "16.1. Автоматизация завершения", style="Heading 2"
    )
    automation = [
        (
            "npm run verify:voice последовательно запускает lint, typecheck, "
            "142 voice-теста, build и проверки VoxEngine."
        ),
        (
            "npm run setup:voximplant создаёт или обновляет приложение, "
            "сценарий, правило и актуальную binding-связь."
        ),
        (
            "npm run smoke:fish-tts проверяет реальный reference voice и "
            "сохраняет WAV для прослушивания."
        ),
        (
            "Отдельный reviewer выполняет независимый gate по race conditions, "
            "security, failure paths и документации."
        ),
    ]
    for text in automation:
        document.add_paragraph(text, style="Plan List")

    document.add_paragraph(
        "16.2. Условия окончательной live-приёмки", style="Heading 2"
    )
    live_conditions = [
        "Выданы Fish, Deepgram, Anthropic и при необходимости AssemblyAI API keys.",
        (
            "Создан разрешённый Fish reference voice; зафиксированы согласие, "
            "commercial rights и AI disclosure."
        ),
        (
            "Развёрнут публичный TLS backend, куплен и привязан номер "
            "Voximplant."
        ),
        (
            "Выполнены PSTN A/B и failure/load tests; подтверждены p50/p95, "
            "entity accuracy, fallback rate, стоимость и отсутствие "
            "противоречий."
        ),
    ]
    for text in live_conditions:
        document.add_paragraph(text, style="Plan List")


def refresh_existing_implementation_status(document: Document) -> None:
    replacements = (
        (
            "Статус: 137 автоматизированных voice-тестов проходят;",
            "Статус: проходят 142 автоматизированных voice-теста;",
        ),
        (
            "127 автоматизированных voice-тестов",
            "142 автоматизированных voice-теста",
        ),
        (
            "135 автоматизированных voice-тестов",
            "142 автоматизированных voice-теста",
        ),
        (
            "137 автоматизированных voice-тестов",
            "142 автоматизированных voice-теста",
        ),
        (
            "127 тестов, lint, typecheck, build, syntax и diff gates.",
            (
                "142 теста, graceful shutdown drain, lint, typecheck, build, "
                "syntax и diff gates."
            ),
        ),
        (
            "135 тестов, graceful shutdown drain, lint, typecheck, build, "
            "syntax и diff gates.",
            (
                "142 теста, graceful shutdown drain, lint, typecheck, build, "
                "syntax и diff gates."
            ),
        ),
        (
            "137 тестов, graceful shutdown drain, lint, typecheck, build, "
            "syntax и diff gates.",
            (
                "142 теста, graceful shutdown drain, lint, typecheck, build, "
                "syntax и diff gates."
            ),
        ),
        ("127 voice-тестов", "142 voice-теста"),
        ("135 voice-тестов", "142 voice-теста"),
        ("137 voice-тестов", "142 voice-теста"),
    )

    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)

    for paragraph in paragraphs:
        text = paragraph.text
        updated = text
        for old, new in replacements:
            updated = updated.replace(old, new)
        if updated != text:
            replace_paragraph(paragraph, updated)


def finalize_appendix_layout(document: Document) -> None:
    status_table = document.tables[-1]
    compact_labels = ["Контур", "Media", "STT", "LLM", "TTS", "Gates", "H100"]
    if len(status_table.rows) == len(compact_labels):
        for row_index, (row, label) in enumerate(
            zip(status_table.rows, compact_labels, strict=True)
        ):
            set_cell_text(
                row.cells[0],
                label,
                bold=row_index == 0,
            )

    paragraphs = document.paragraphs
    combined_acceptance = (
        "Выполнены PSTN A/B и failure/load tests; подтверждены p50/p95, "
        "entity accuracy, fallback rate, стоимость и отсутствие противоречий."
    )
    for paragraph in paragraphs:
        if paragraph.text.strip() == (
            "Выполнены PSTN A/B, barge-in, provider outage, длинный звонок и "
            "параллельная нагрузка."
        ):
            replace_paragraph(paragraph, combined_acceptance)
        elif paragraph.text.strip() == (
            "Подтверждены p50/p95 latency, entity accuracy, fallback rate, "
            "стоимость и отсутствие противоречий."
        ):
            paragraph._element.getparent().remove(paragraph._element)

    paragraphs = document.paragraphs
    for index, paragraph in enumerate(paragraphs):
        if paragraph.text.strip() != APPENDIX_HEADING or index == 0:
            continue
        previous = paragraphs[index - 1]
        if (
            not previous.text.strip()
            and previous._element.xpath(".//w:br[@w:type='page']")
        ):
            previous._element.getparent().remove(previous._element)
        break

    paragraphs = document.paragraphs
    for index, paragraph in enumerate(paragraphs):
        if paragraph.text.strip() != "16.3. Итог":
            continue
        if index + 1 < len(paragraphs):
            next_paragraph = paragraphs[index + 1]
            next_paragraph._element.getparent().remove(
                next_paragraph._element
            )
        paragraph._element.getparent().remove(paragraph._element)
        break


def main() -> None:
    document = Document(DOCUMENT_PATH)
    replace_existing_status(document)
    update_tables(document)
    update_local_file_index(document)
    append_implementation_status(document)
    refresh_existing_implementation_status(document)
    finalize_appendix_layout(document)
    document.save(DOCUMENT_PATH)
    print(DOCUMENT_PATH)


if __name__ == "__main__":
    main()
